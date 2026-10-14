"""Репозиторий для выполнения операций с товарами, снимками цен, пользователями и настройками."""

from datetime import datetime, timedelta
from decimal import Decimal
import io
import re
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import desc, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from config.settings import settings
from database.models import PriceHistory, Product, SystemSetting, User
from schemas.tracker import ParsedProduct


class ProductRepository:
    """Интерфейс доступа к данным товаров с изоляцией по пользователям."""

    def __init__(self, session: AsyncSession) -> None:
        self._session = session

    # --- Управление пользователями и правами доступа ---

    async def sync_user_profile(
        self,
        chat_id: int,
        username: Optional[str] = None,
        full_name: Optional[str] = None,
    ) -> User:
        """Синхронизация профиля пользователя (юзернейм и имя) при любой активности."""
        user = await self.get_user(chat_id)
        admin_id = settings.ADMIN_CHAT_ID or settings.TELEGRAM_ALERT_CHAT_ID
        is_admin = bool(admin_id and chat_id == admin_id)

        if not user:
            user = User(
                telegram_chat_id=chat_id,
                username=username,
                full_name=full_name,
                is_active=is_admin,
                max_products=10,
            )
            self._session.add(user)
        else:
            if username:
                user.username = username
            if full_name:
                user.full_name = full_name
            if is_admin:
                user.is_active = True
        await self._session.flush()
        return user

    async def is_user_allowed(self, chat_id: int) -> bool:
        """Проверка прав доступа: админ или активный пользователь из базы."""
        admin_id = settings.ADMIN_CHAT_ID or settings.TELEGRAM_ALERT_CHAT_ID
        if admin_id and chat_id == admin_id:
            return True

        stmt = select(User).where(
            User.telegram_chat_id == chat_id, User.is_active.is_(True)
        )
        result = await self._session.execute(stmt)
        return result.scalar_one_or_none() is not None

    async def get_user(self, chat_id: int) -> Optional[User]:
        """Получение профиля пользователя по chat_id."""
        stmt = select(User).where(User.telegram_chat_id == chat_id)
        result = await self._session.execute(stmt)
        return result.scalar_one_or_none()

    async def grant_user_access(
        self,
        chat_id: int,
        username: Optional[str] = None,
        full_name: Optional[str] = None,
    ) -> User:
        """Выдача или активация доступа пользователю."""
        user = await self.get_user(chat_id)
        if not user:
            user = User(
                telegram_chat_id=chat_id,
                username=username,
                full_name=full_name,
                is_active=True,
                max_products=10,
            )
            self._session.add(user)
        else:
            user.is_active = True
            if username:
                user.username = username
            if full_name:
                user.full_name = full_name
        await self._session.flush()
        return user

    async def revoke_user_access(self, chat_id: int) -> bool:
        """Отзыв доступа у пользователя."""
        user = await self.get_user(chat_id)
        if not user:
            return False
        user.is_active = False
        await self._session.flush()
        return True

    async def update_user_max_products(self, chat_id: int, max_products: int) -> bool:
        """Изменение лимита товаров пользователя (смена тарифного плана)."""
        user = await self.get_user(chat_id)
        if not user:
            return False
        user.max_products = max_products
        await self._session.flush()
        return True

    async def link_team_chat(self, chat_id: int, team_chat_id: Optional[int]) -> bool:
        """Привязка или отключение рабочего чата команды селлера."""
        user = await self.get_user(chat_id)
        if not user:
            return False
        user.team_chat_id = team_chat_id
        await self._session.flush()
        return True

    async def get_user_team_chat(self, chat_id: int) -> Optional[int]:
        """Получение идентификатора командного чата селлера."""
        user = await self.get_user(chat_id)
        return user.team_chat_id if user and user.is_active else None

    async def get_all_users(self) -> List[User]:
        """Получение списка всех зарегистрированных пользователей."""
        stmt = select(User).order_by(desc(User.created_at))
        result = await self._session.execute(stmt)
        return list(result.scalars().all())

    async def get_users_overview(self) -> List[Tuple[User, int]]:
        """Получение списка пользователей с количеством их активных товаров."""
        users = await self.get_all_users()
        overview: List[Tuple[User, int]] = []
        for u in users:
            cnt = await self.count_user_products(u.telegram_chat_id)
            overview.append((u, cnt))
        return overview

    async def count_user_products(self, chat_id: int) -> int:
        """Подсчет количества активных товаров у пользователя."""
        stmt = select(func.count(Product.id)).where(
            Product.telegram_chat_id == chat_id, Product.is_active.is_(True)
        )
        result = await self._session.execute(stmt)
        return result.scalar_one() or 0

    # --- Системная аналитика для админа ---

    async def get_system_metrics(self) -> Dict[str, Any]:
        """Сбор сводных коммерческих метрик системы."""
        total_users = (
            await self._session.execute(select(func.count(User.id)))
        ).scalar_one() or 0
        active_users = (
            await self._session.execute(
                select(func.count(User.id)).where(User.is_active.is_(True))
            )
        ).scalar_one() or 0
        total_products = (
            await self._session.execute(
                select(func.count(Product.id)).where(Product.is_active.is_(True))
            )
        ).scalar_one() or 0
        total_snapshots = (
            await self._session.execute(select(func.count(PriceHistory.id)))
        ).scalar_one() or 0

        return {
            "total_users": total_users,
            "active_users": active_users,
            "total_products": total_products,
            "total_snapshots": total_snapshots,
        }

    # --- Полная очистка БД ---

    async def wipe_database(self) -> None:
        """Полная очистка всех таблиц базы данных и сброс структуры."""
        from database.models import Base
        from database.session import engine

        async with engine.begin() as conn:
            await conn.run_sync(Base.metadata.drop_all)
            await conn.run_sync(Base.metadata.create_all)

    # --- Генерация Excel-отчета ---

    async def export_product_history_excel(
        self, chat_id: int, sku: str
    ) -> Tuple[Optional[bytes], str]:
        """Генерация чистого Excel-отчета с автофильтром, сортировкой и подбором ширины."""
        product = await self.get_by_sku(chat_id=chat_id, sku=sku)
        if not product:
            return None, f"Ozon_{sku}.xlsx"

        clean_title = re.sub(r'[\\/*?:"<>|]', "", product.title)
        clean_title = re.sub(r"\s+", "_", clean_title).strip("_")[:30]
        filename = f"Ozon_{sku}_{clean_title}.xlsx"

        stmt = (
            select(PriceHistory)
            .where(PriceHistory.product_id == product.id)
            .order_by(desc(PriceHistory.recorded_at))
        )
        result = await self._session.execute(stmt)
        history = list(result.scalars().all())

        if not history:
            return None, filename

        wb = Workbook()
        active_sheet = wb.active
        if isinstance(active_sheet, Worksheet):
            ws = active_sheet
            ws.title = "История цен"
        else:
            ws = wb.create_sheet(title="История цен")

        ws.views.sheetView[0].showGridLines = True

        font_brand = Font(name="Segoe UI", size=13, bold=True, color="1F4E78")
        font_meta_lbl = Font(name="Segoe UI", size=9, bold=True, color="595959")
        font_meta_val = Font(name="Segoe UI", size=9, color="262626")
        font_th = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        font_td = Font(name="Segoe UI", size=9)
        font_bold = Font(name="Segoe UI", size=9, bold=True)
        font_green = Font(name="Segoe UI", size=9, bold=True, color="1E7E34")
        font_red = Font(name="Segoe UI", size=9, bold=True, color="BD2130")

        fill_header = PatternFill(
            start_color="1F4E78", end_color="1F4E78", fill_type="solid"
        )
        fill_zebra = PatternFill(
            start_color="F8FAFC", end_color="F8FAFC", fill_type="solid"
        )

        align_center = Alignment(horizontal="center", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")

        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
        )

        c_title = ws["A1"]
        if isinstance(c_title, Cell):
            c_title.value = "ОТЧЕТ ПО МОНИТОРИНГУ ЦЕН OZON"
            c_title.font = font_brand

        target_str = (
            f"{product.target_price:,.0f} ₽".replace(",", " ")
            if product.target_price
            else "Не задана"
        )
        meta_info = [
            ("Товар:", product.title),
            ("Артикул SKU:", str(product.sku)),
            ("Порог демпинга:", f"{product.dumping_threshold}%"),
            ("Минимальная цена (себестоимость):", target_str),
            ("Всего замеров в базе:", f"{len(history)}"),
        ]

        for idx, (label, val) in enumerate(meta_info, start=2):
            lbl_cell = ws.cell(row=idx, column=1, value=label)
            lbl_cell.font = font_meta_lbl
            lbl_cell.alignment = align_left
            val_cell = ws.cell(row=idx, column=2, value=val)
            val_cell.font = font_meta_val
            val_cell.alignment = align_left

        headers = [
            "Дата (МСК)",
            "Время (МСК)",
            "Обычная цена",
            "По Ozon Карте",
            "Скидка по карте",
            "Старая цена",
            "Статус",
        ]

        header_row = 7
        for col_idx, h_text in enumerate(headers, start=1):
            cell = ws.cell(row=header_row, column=col_idx, value=h_text)
            cell.font = font_th
            cell.fill = fill_header
            cell.alignment = align_center

        start_row = 8
        cur_row = start_row
        for row_offset, snap in enumerate(history):
            cur_row = start_row + row_offset
            msk_time = snap.recorded_at + timedelta(hours=3)

            date_str = msk_time.strftime("%d.%m.%Y")
            time_str = msk_time.strftime("%H:%M")

            current_p = float(snap.current_price)
            card_p = float(snap.card_price) if snap.card_price else None
            card_diff = (
                (current_p - card_p) if (card_p and current_p > card_p) else None
            )
            orig_p = float(snap.original_price) if snap.original_price else None
            stock_str = "В наличии" if snap.is_in_stock else "Нет на складе"

            row_data = [
                (date_str, align_center, font_td, None),
                (time_str, align_center, font_td, None),
                (current_p, align_right, font_bold, "#,##0 ₽"),
                (
                    card_p if card_p else "-",
                    align_right,
                    font_bold,
                    "#,##0 ₽" if card_p else None,
                ),
                (
                    card_diff if card_diff else "-",
                    align_right,
                    font_td,
                    "#,##0 ₽" if card_diff else None,
                ),
                (
                    orig_p if orig_p else "-",
                    align_right,
                    font_td,
                    "#,##0 ₽" if orig_p else None,
                ),
                (
                    stock_str,
                    align_center,
                    font_green if snap.is_in_stock else font_red,
                    None,
                ),
            ]

            is_zebra = row_offset % 2 == 1

            for col_idx, (val, alignment, font, num_fmt) in enumerate(
                row_data, start=1
            ):
                cell = ws.cell(row=cur_row, column=col_idx, value=val)
                cell.font = font
                cell.alignment = alignment
                cell.border = thin_border
                if num_fmt:
                    cell.number_format = num_fmt
                if is_zebra:
                    cell.fill = fill_zebra

        ws.freeze_panes = "A8"
        ws.auto_filter.ref = f"A7:G{cur_row}"

        for col in list(ws.columns)[:7]:
            max_len = 0
            first_cell = col[0]
            if not isinstance(first_cell, Cell) or first_cell.column is None:
                continue
            col_letter = get_column_letter(first_cell.column)

            for cell in col:
                if not isinstance(cell, Cell):
                    continue
                if cell.row is not None and cell.row < 7:
                    continue
                if cell.value is not None:
                    val_str = str(cell.value)
                    if len(val_str) > max_len:
                        max_len = len(val_str)

            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue(), filename

    # --- Системные настройки ---

    async def get_setting(self, key: str, default: str) -> str:
        """Получение системной настройки из БД."""
        stmt = select(SystemSetting).where(SystemSetting.key == key)
        result = await self._session.execute(stmt)
        item = result.scalar_one_or_none()
        return item.value if item else default

    async def set_setting(self, key: str, value: str) -> None:
        """Сохранение системной настройки в БД."""
        stmt = select(SystemSetting).where(SystemSetting.key == key)
        result = await self._session.execute(stmt)
        item = result.scalar_one_or_none()
        if not item:
            item = SystemSetting(key=key, value=value)
            self._session.add(item)
        else:
            item.value = value
        await self._session.flush()

    # --- Управление товарами ---

    async def get_by_sku(self, chat_id: int, sku: str) -> Optional[Product]:
        """Получение товара конкретного пользователя по артикулу SKU."""
        stmt = select(Product).where(
            Product.telegram_chat_id == chat_id,
            Product.sku == str(sku),
        )
        result = await self._session.execute(stmt)
        return result.scalar_one_or_none()

    async def delete_by_sku(self, chat_id: int, sku: str) -> bool:
        """Удаление товара конкретного пользователя по артикулу SKU."""
        product = await self.get_by_sku(chat_id=chat_id, sku=sku)
        if not product:
            return False
        await self._session.delete(product)
        await self._session.flush()
        return True

    async def update_product_threshold(
        self, chat_id: int, sku: str, threshold: float
    ) -> bool:
        """Обновление порога демпинга для товара конкретного пользователя."""
        product = await self.get_by_sku(chat_id=chat_id, sku=sku)
        if not product:
            return False
        product.dumping_threshold = threshold
        await self._session.flush()
        return True

    async def update_product_target_price(
        self, chat_id: int, sku: str, target_price: Optional[Decimal]
    ) -> bool:
        """Обновление порога минимальной цены (себестоимости) в рублях."""
        product = await self.get_by_sku(chat_id=chat_id, sku=sku)
        if not product:
            return False
        product.target_price = target_price
        await self._session.flush()
        return True

    async def update_product_interval(
        self, chat_id: int, sku: str, minutes: int
    ) -> bool:
        """Обновление интервала проверки для товара конкретного пользователя."""
        product = await self.get_by_sku(chat_id=chat_id, sku=sku)
        if not product:
            return False
        product.check_interval_minutes = minutes
        await self._session.flush()
        return True

    async def update_last_checked(self, product_id: int, timestamp: datetime) -> None:
        """Обновление времени последнего замера по первичному ключу товара."""
        stmt = select(Product).where(Product.id == product_id)
        result = await self._session.execute(stmt)
        product = result.scalar_one_or_none()
        if product:
            product.last_checked_at = timestamp
            await self._session.flush()

    async def get_active_products(self, chat_id: Optional[int] = None) -> List[Product]:
        """Получение активных товаров с опциональным фильтром по пользователю."""
        stmt = select(Product).where(Product.is_active.is_(True))
        if chat_id is not None:
            stmt = stmt.where(Product.telegram_chat_id == chat_id)
        result = await self._session.execute(stmt)
        return list(result.scalars().all())

    async def get_due_products(self, limit: int = 25) -> List[Product]:
        """Получение очереди товаров, готовых к проверке, напрямую через SQL-фильтрацию."""
        next_check_expr = func.datetime(
            Product.last_checked_at,
            func.concat("+", Product.check_interval_minutes, " minutes"),
        )
        stmt = (
            select(Product)
            .where(
                Product.is_active.is_(True),
                next_check_expr <= func.datetime("now"),
            )
            .order_by(Product.last_checked_at.asc())
            .limit(limit)
        )
        result = await self._session.execute(stmt)
        return list(result.scalars().all())

    async def add_or_update_product(
        self,
        parsed: ParsedProduct,
        chat_id: int,
        dumping_threshold: Optional[float] = None,
        check_interval_minutes: Optional[int] = None,
    ) -> Product:
        """Создание карточки товара с привязкой к chat_id или ее обновление."""
        product = await self.get_by_sku(chat_id=chat_id, sku=parsed.sku)
        if not product:
            product = Product(
                telegram_chat_id=chat_id,
                sku=parsed.sku,
                title=parsed.title,
                url=str(parsed.url),
                dumping_threshold=(
                    dumping_threshold if dumping_threshold is not None else 5.0
                ),
                check_interval_minutes=(
                    check_interval_minutes if check_interval_minutes is not None else 30
                ),
                last_checked_at=datetime.utcnow(),
                is_active=True,
            )
            self._session.add(product)
            await self._session.flush()
        else:
            product.title = parsed.title
            product.url = str(parsed.url)
            product.last_checked_at = datetime.utcnow()
            if dumping_threshold is not None:
                product.dumping_threshold = dumping_threshold
            if check_interval_minutes is not None:
                product.check_interval_minutes = check_interval_minutes
            await self._session.flush()

        return product

    async def save_price_snapshot(
        self, product_id: int, parsed: ParsedProduct
    ) -> PriceHistory:
        """Сохранение нового замера цен и остатка в историю."""
        snapshot = PriceHistory(
            product_id=product_id,
            current_price=parsed.current_price,
            card_price=parsed.card_price,
            original_price=parsed.original_price,
            is_in_stock=parsed.is_in_stock,
            stock_count=parsed.stock_count,
            recorded_at=parsed.parsed_at,
        )
        self._session.add(snapshot)
        await self._session.flush()
        return snapshot

    async def get_last_two_snapshots(
        self, product_id: int
    ) -> Tuple[Optional[PriceHistory], Optional[PriceHistory]]:
        """Получение двух последних снимков цен (текущий и предыдущий)."""
        stmt = (
            select(PriceHistory)
            .where(PriceHistory.product_id == product_id)
            .order_by(desc(PriceHistory.recorded_at))
            .limit(2)
        )
        result = await self._session.execute(stmt)
        snapshots = list(result.scalars().all())

        latest = snapshots[0] if len(snapshots) > 0 else None
        previous = snapshots[1] if len(snapshots) > 1 else None
        return latest, previous
